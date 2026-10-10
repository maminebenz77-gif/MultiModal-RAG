"""generate_document: the last node in the graph, run only after a
human has approved the final answers in human_review. Deterministic --
no LLM call -- because by the time a run reaches here, every answer
has already been judged (by validate_answer or a human); export only
lays out text that is already final, it never generates or changes
any of it.

There is no existing docgen template FILE to fill into --
Configuration.template is a display name a person typed in free text
(nodes/configuration.py), not a path to a .pptx/.docx on disk. So
"export" here means building a fresh presentation/document from
python-pptx's/python-docx's own blank defaults and filling it with one
slide/section per question, not merging into a pre-built template.
That's a real simplification worth naming: a future phase could add
real template files (logos, fixed slide masters) once someone actually
needs one, rather than speculatively building template-lookup logic
with nothing yet to look up.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from pptx import Presentation
from pptx.presentation import Presentation as PptxPresentation

from ..state import Answer, Configuration, Question

DEFAULT_OUTPUT_DIR = Path("data/docgen_outputs")
_SLUG_RE = re.compile(r"[^a-z0-9]+")

_TITLE_AND_CONTENT_LAYOUT = 1


def _slugify(name: str) -> str:
    slug = _SLUG_RE.sub("-", name.strip().lower()).strip("-")
    return slug or "docgen-report"


def _ordered_answers(
    questions: list[Question], answers: dict[str, Answer]
) -> list[tuple[Question, Answer]]:
    """Export follows the ORIGINAL question order, not answers'
    arbitrary dict order -- the same reason _format_review_summary
    (nodes/review.py) iterates questions rather than answers."""
    pairs = []
    for question in questions:
        answer = answers.get(question["id"])
        if answer is not None:
            pairs.append((question, answer))
    return pairs


def _build_pptx(
    title: str, questions: list[Question], answers: dict[str, Answer]
) -> PptxPresentation:
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    body_layout = presentation.slide_layouts[_TITLE_AND_CONTENT_LAYOUT]
    for question, answer in _ordered_answers(questions, answers):
        slide = presentation.slides.add_slide(body_layout)
        slide.shapes.title.text = question["text"]
        slide.placeholders[1].text = answer["text"]
    return presentation


def _build_docx(
    title: str, questions: list[Question], answers: dict[str, Answer]
) -> DocxDocument:
    document = Document()
    document.add_heading(title, level=0)
    for question, answer in _ordered_answers(questions, answers):
        document.add_heading(question["text"], level=1)
        document.add_paragraph(answer["text"])
    return document


def generate_document(
    questions: list[Question],
    answers: dict[str, Answer],
    configuration: Configuration,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
) -> Path:
    title = configuration["template"] or "Generated Report"
    output_dir.mkdir(parents=True, exist_ok=True)
    slug = _slugify(configuration["template"] or "docgen-report")
    if configuration["format"] == "pptx":
        path = output_dir / f"{slug}.pptx"
        _build_pptx(title, questions, answers).save(str(path))
    else:
        path = output_dir / f"{slug}.docx"
        _build_docx(title, questions, answers).save(str(path))
    return path
