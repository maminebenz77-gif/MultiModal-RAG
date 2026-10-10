from pathlib import Path

from docx import Document
from pptx import Presentation

from multimodal_rag.docgen.nodes.export import generate_document
from multimodal_rag.docgen.state import Answer, Question


def _question(id_: str, text: str) -> Question:
    return {"id": id_, "text": text, "sources_required": ["task_docs"], "status": "answered"}


def _answer(text: str) -> Answer:
    return {"text": text, "attempts": [], "accepted_by": "validation"}


def test_generate_document_writes_a_pptx_with_one_slide_per_question(tmp_path: Path) -> None:
    questions = [_question("q1", "What was the latency?"), _question("q2", "What was the cost?")]
    answers = {"q1": _answer("120ms."), "q2": _answer("$5.")}
    configuration = {
        "template": "My Report",
        "format": "pptx",
        "confirmed": True,
        "max_retries": 3,
    }

    path = generate_document(questions, answers, configuration, tmp_path)

    assert path == tmp_path / "my-report.pptx"
    presentation = Presentation(path)
    assert len(presentation.slides) == 3  # title slide + one per question
    assert presentation.slides[0].shapes.title.text == "My Report"
    assert presentation.slides[1].shapes.title.text == "What was the latency?"
    assert presentation.slides[1].placeholders[1].text == "120ms."
    assert presentation.slides[2].shapes.title.text == "What was the cost?"
    assert presentation.slides[2].placeholders[1].text == "$5."


def test_generate_document_writes_a_docx_with_one_heading_per_question(tmp_path: Path) -> None:
    questions = [_question("q1", "What was the latency?")]
    answers = {"q1": _answer("120ms.")}
    configuration = {"template": "", "format": "docx", "confirmed": True, "max_retries": 3}

    path = generate_document(questions, answers, configuration, tmp_path)

    assert path == tmp_path / "docgen-report.docx"
    document = Document(path)
    paragraphs = [p.text for p in document.paragraphs]
    assert "Generated Report" in paragraphs
    assert "What was the latency?" in paragraphs
    assert "120ms." in paragraphs


def test_generate_document_skips_a_question_with_no_accepted_answer(tmp_path: Path) -> None:
    questions = [_question("q1", "Answered?"), _question("q2", "Never answered?")]
    answers = {"q1": _answer("Yes.")}
    configuration = {"template": "", "format": "docx", "confirmed": True, "max_retries": 3}

    path = generate_document(questions, answers, configuration, tmp_path)

    document = Document(path)
    paragraphs = [p.text for p in document.paragraphs]
    assert "Never answered?" not in paragraphs
