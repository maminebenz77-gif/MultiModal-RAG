import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference

from multimodal_rag.ingestion import excel_charts
from multimodal_rag.ingestion.excel_charts import charts_for_worksheet, describe_chart


def _write_bar_chart_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(["Region", "Revenue"])
    worksheet.append(["North", 120])
    worksheet.append(["South", 340])
    worksheet.append(["East", 275])

    chart = BarChart()
    chart.type = "col"
    chart.title = "Q3 Revenue by Region"
    data = Reference(worksheet, min_col=2, min_row=1, max_row=4)
    categories = Reference(worksheet, min_col=1, min_row=2, max_row=4)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    worksheet.add_chart(chart, "D2")

    workbook.save(path)
    return path


def _rewrite_chart_target_as_package_relative(path: Path) -> None:
    """openpyxl's own writer always emits a package-ABSOLUTE relationship
    target ("/xl/charts/chart1.xml"). Real Excel (and LibreOffice) write
    the spec-correct form instead: relative to the SOURCE part's own
    directory ("../charts/chart1.xml" from "xl/drawings/") -- see
    excel_charts._resolve_part_path's docstring for why that distinction
    caused a real KeyError this project hit. Every other test in this
    file builds its workbook with openpyxl and would never exercise
    that path, so this rewrites one .rels part byte-for-byte after
    saving, standing in for a real Excel-produced file without checking
    one into the repo.
    """
    with zipfile.ZipFile(path) as zf:
        entries = {name: zf.read(name) for name in zf.namelist()}
    rels_name = "xl/drawings/_rels/drawing1.xml.rels"
    entries[rels_name] = entries[rels_name].replace(
        b'Target="/xl/charts/chart1.xml"', b'Target="../charts/chart1.xml"'
    )
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in entries.items():
            zf.writestr(name, data)


@pytest.fixture
def _stub_llm(monkeypatch: pytest.MonkeyPatch) -> None:
    # Not autouse -- test_summarize_returns_none_when_* need the REAL
    # _summarize. Every describe_chart test that only cares about correct
    # XML extraction/resolution (not the LLM call itself) requests this
    # explicitly to get the raw structured description back instead.
    monkeypatch.setattr(excel_charts, "_summarize", lambda description: description)


def test_charts_for_worksheet_finds_the_chart_on_its_own_sheet(tmp_path: Path) -> None:
    path = _write_bar_chart_workbook(tmp_path / "wb.xlsx")
    workbook = openpyxl.load_workbook(path, data_only=True)

    chart_paths = charts_for_worksheet(path, workbook["Sheet1"])

    assert chart_paths == ["xl/charts/chart1.xml"]


def test_charts_for_worksheet_resolves_a_package_relative_chart_target(
    tmp_path: Path,
) -> None:
    # Regression for a real KeyError hit ingesting an actual Excel-saved
    # workbook: real Excel writes the chart's relationship Target
    # relative to xl/drawings/ ("../charts/chart1.xml"), not as a
    # package-absolute path. Resolving it wrong doesn't fail loudly on
    # a mismatched-but-plausible path -- it produces a string ("../charts
    # /chart1.xml") that isn't ANY real entry in the archive, so
    # zipfile.read raises KeyError deep inside describe_chart.
    path = _write_bar_chart_workbook(tmp_path / "wb.xlsx")
    _rewrite_chart_target_as_package_relative(path)
    workbook = openpyxl.load_workbook(path, data_only=True)

    chart_paths = charts_for_worksheet(path, workbook["Sheet1"])

    assert chart_paths == ["xl/charts/chart1.xml"]


def test_describe_chart_does_not_raise_on_a_package_relative_target(
    tmp_path: Path, _stub_llm: None
) -> None:
    # Same regression, end to end through describe_chart (what
    # excel.py's parse_excel actually calls) -- proves the resolved
    # path is also a real, readable archive entry, not just a string
    # that happens to equal the expected one.
    path = _write_bar_chart_workbook(tmp_path / "wb.xlsx")
    _rewrite_chart_target_as_package_relative(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook["Sheet1"]
    [chart_path] = charts_for_worksheet(path, worksheet)

    description = describe_chart(path, chart_path, worksheet)

    assert description is not None
    assert "North: 120" in description


def test_charts_for_worksheet_empty_when_sheet_has_no_chart(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "Sheet1"
    workbook.active.append(["A"])
    path = tmp_path / "wb.xlsx"
    workbook.save(path)
    reopened = openpyxl.load_workbook(path)

    assert charts_for_worksheet(path, reopened["Sheet1"]) == []


def test_describe_chart_extracts_type_title_and_resolved_series(
    tmp_path: Path, _stub_llm: None
) -> None:
    path = _write_bar_chart_workbook(tmp_path / "wb.xlsx")
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook["Sheet1"]
    [chart_path] = charts_for_worksheet(path, worksheet)

    description = describe_chart(path, chart_path, worksheet)

    assert description is not None
    assert "barChart" in description
    assert "Q3 Revenue by Region" in description
    assert "North: 120" in description
    assert "South: 340" in description
    assert "East: 275" in description


def test_llm_is_used_to_summarize(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(excel_charts, "_summarize", lambda description: "a chart summary")
    path = _write_bar_chart_workbook(tmp_path / "wb.xlsx")
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook["Sheet1"]
    [chart_path] = charts_for_worksheet(path, worksheet)

    assert describe_chart(path, chart_path, worksheet) == "a chart summary"


def test_describe_chart_returns_none_when_summarize_returns_none(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(excel_charts, "_summarize", lambda description: None)
    path = _write_bar_chart_workbook(tmp_path / "wb.xlsx")
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook["Sheet1"]
    [chart_path] = charts_for_worksheet(path, worksheet)

    assert describe_chart(path, chart_path, worksheet) is None


def test_summarize_returns_none_when_llm_unconfigured(monkeypatch: pytest.MonkeyPatch) -> None:
    def _raise() -> None:
        raise NotImplementedError("no provider configured")

    monkeypatch.setattr("multimodal_rag.ingestion.excel_charts.get_llm", _raise)
    assert excel_charts._summarize("some description") is None


def test_summarize_returns_none_when_llm_call_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    class _FailingLLM:
        def generate(self, messages: list[dict[str, str]]) -> str:
            raise RuntimeError("provider unavailable")

    monkeypatch.setattr("multimodal_rag.ingestion.excel_charts.get_llm", lambda: _FailingLLM())
    assert excel_charts._summarize("some description") is None
